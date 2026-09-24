"""Raw File/Event -> Type Detection -> Extraction/OCR.

Every extractor returns a list of TextUnit objects. A TextUnit's `location_ref`
is the provenance anchor (page / row / paragraph / event id) required by the
relationship-extraction and entity-resolution stages downstream - it must
never be dropped, since it is what lets an investigator trace a finding back
to the exact spot in the source document.
"""

import io
import json
from dataclasses import dataclass
from typing import List

import fitz  # PyMuPDF
import pandas as pd
import pytesseract
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image


@dataclass
class TextUnit:
    text: str
    location_ref: str


def extract_pdf(content: bytes) -> List[TextUnit]:
    units: List[TextUnit] = []
    with fitz.open(stream=content, filetype="pdf") as doc:
        for page_index, page in enumerate(doc):
            text = page.get_text().strip()
            if text:
                units.append(TextUnit(text=text, location_ref=f"page:{page_index + 1}"))
    return units


def extract_docx(content: bytes) -> List[TextUnit]:
    doc = DocxDocument(io.BytesIO(content))
    units: List[TextUnit] = []
    for para_index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            units.append(TextUnit(text=text, location_ref=f"paragraph:{para_index + 1}"))
    return units


def extract_csv(content: bytes) -> List[TextUnit]:
    df = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False)
    units: List[TextUnit] = []
    for row_index, row in df.iterrows():
        text = " | ".join(f"{col}: {val}" for col, val in row.items() if str(val).strip())
        if text:
            units.append(TextUnit(text=text, location_ref=f"row:{row_index + 1}"))
    return units


def extract_xlsx(content: bytes) -> List[TextUnit]:
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    units: List[TextUnit] = []
    for sheet in workbook.worksheets:
        headers = [str(c.value) if c.value is not None else "" for c in next(sheet.iter_rows(max_row=1), [])]
        for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            values = [str(c.value) if c.value is not None else "" for c in row]
            text = " | ".join(
                f"{headers[i] if i < len(headers) else i}: {v}" for i, v in enumerate(values) if v.strip()
            )
            if text:
                units.append(TextUnit(text=text, location_ref=f"sheet:{sheet.title}|row:{row_index}"))
    return units


def extract_image_ocr(content: bytes) -> List[TextUnit]:
    image = Image.open(io.BytesIO(content))
    text = pytesseract.image_to_string(image).strip()
    if not text:
        return []
    return [TextUnit(text=text, location_ref="ocr:whole-image")]


def extract_json(content: bytes) -> List[TextUnit]:
    data = json.loads(content)
    units: List[TextUnit] = []
    records = data if isinstance(data, list) else [data]
    for index, record in enumerate(records):
        text = json.dumps(record, ensure_ascii=False)
        event_id = record.get("event_id") or record.get("id") if isinstance(record, dict) else None
        location_ref = f"event:{event_id}" if event_id else f"json_item:{index}"
        units.append(TextUnit(text=text, location_ref=location_ref))
    return units


def extract_txt(content: bytes) -> List[TextUnit]:
    text = content.decode("utf-8", errors="replace").strip()
    if not text:
        return []
    return [TextUnit(text=text, location_ref="whole-document")]


_MIME_DISPATCH = {
    "application/pdf": extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extract_docx,
    "text/csv": extract_csv,
    "application/vnd.ms-excel": extract_xlsx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": extract_xlsx,
    "image/png": extract_image_ocr,
    "image/jpeg": extract_image_ocr,
    "image/tiff": extract_image_ocr,
    "application/json": extract_json,
    "text/plain": extract_txt,
}


def extract_text_units(mime_type: str, content: bytes) -> List[TextUnit]:
    extractor = _MIME_DISPATCH.get(mime_type)
    if extractor is None:
        raise ValueError(f"No extractor registered for mime type: {mime_type}")
    return extractor(content)
